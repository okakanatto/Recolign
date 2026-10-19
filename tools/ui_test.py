#!/usr/bin/env python3
"""UI tests of the desktop app's screens, run in Chromium (the engine of
WebView2) against the development harness.

    python tools/ui_test.py --harness target/release/cdr-harness [--screens DIR]

Covers the setup flow, the comparison screen (counts, navigation, orders,
filters, sort, copy, go to key/line, detail, hidden columns, layouts,
keyboard) and the key-problem screen, using samples/. With --screens,
screenshots are taken at 1366x768 and 1920x1080 with 100/125/150 % scaling.
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import tempfile
import traceback
from pathlib import Path

from playwright.sync_api import Page, sync_playwright

ROOT = Path(__file__).resolve().parent.parent
CHROMIUM = os.environ.get('CHROMIUM_EXECUTABLE', '/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
SAMPLES = ROOT / 'samples'
EXPECTED = json.loads((SAMPLES / 'expected.json').read_text(encoding='utf-8'))

WAIT_IDLE = '''() => new Promise((res) => {
  const chk = () => {
    const s = window.__cdr && window.__cdr.state;
    if (!s || s.busy) return requestAnimationFrame(chk);
    const L = s.panes.left, R = s.panes.right;
    if (L.model && R.model && L.complete && R.complete && !L.model.inflight && !R.model.inflight) res(true);
    else requestAnimationFrame(chk);
  };
  requestAnimationFrame(chk);
})'''

API = '''async ([cmd, args]) => {
  const t = new URLSearchParams(location.hash.slice(1)).get('t');
  const r = await fetch('api/' + cmd, {method: 'POST', headers: {'Content-Type': 'application/json', 'X-CDR-Token': t},
    body: JSON.stringify(args)});
  return r.json();
}'''


class Harness:
    def __init__(self, exe: Path, args: list[str]):
        self.p = subprocess.Popen([str(exe), '--ui', str(ROOT / 'app' / 'dist'), *args], stdout=subprocess.PIPE, text=True)
        self.url = self.p.stdout.readline().strip()

    def close(self):
        self.p.terminate()
        self.p.wait()


results: list[tuple[str, bool, str]] = []
# Paths typed into the harness's path prompts, in order.
answers: list[str] = []


def check(name: str, cond: bool, detail: str = '') -> None:
    results.append((name, bool(cond), detail))
    print(('ok   ' if cond else 'FAIL ') + name + (f' ({detail})' if detail and not cond else ''))


def idle(page: Page) -> None:
    page.evaluate(WAIT_IDLE)
    page.wait_for_timeout(30)


def api(page: Page, cmd: str, args: dict | None = None):
    return page.evaluate(API, [cmd, args or {}])


def sel(page: Page):
    return page.evaluate('__cdr.state.sel && {unit: __cdr.state.sel.unit, rows: __cdr.state.sel.rows}')


def wait_sel(page: Page) -> dict:
    page.wait_for_function('__cdr.state.sel && __cdr.state.sel.unit >= 0')
    idle(page)
    return sel(page)


def key_of(page: Page, unit: int) -> list[str]:
    return api(page, 'detail', {'unit': unit})['key']


def total(page: Page, side='left') -> int:
    return page.evaluate(f"__cdr.state.panes.{side}.model.total")


def on_dialog(d) -> None:
    if answers:
        d.accept(answers.pop(0))
    else:
        d.dismiss()


def choose_key(page: Page, name: str) -> None:
    """Makes `name` the key chosen by hand (a guessed key is pressed already:
    pressing it takes it off, pressing again chooses it)."""
    chip = page.locator(f'.chip[data-name="{name}"]').first
    if chip.get_attribute('aria-pressed') == 'true':
        chip.click()
    chip.click()


def setup_flow(page: Page, url: str) -> None:
    page.goto(url)
    page.wait_for_selector('main.setup')
    check('setup: compare button disabled before files', page.is_disabled('#compareBtn'))
    answers.append(str(SAMPLES / 'before.csv'))
    page.locator('.card').nth(0).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .pv')
    answers.append(str(SAMPLES / 'after.csv'))
    page.locator('.card').nth(1).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    enc = page.locator('.card[data-side=left] select').first.input_value()
    check('setup: encoding suggested as UTF-8', enc == 'utf8', enc)
    chips = page.locator('.block .chips').first.locator('.chip')
    check('setup: common columns offered as keys', chips.count() == 10, str(chips.count()))
    pressed = lambda: [c.get_attribute('data-name') for c in chips.all() if c.get_attribute('aria-pressed') == 'true']
    check('setup: the key is guessed from the files', pressed() == ['社員番号'] and '推定' in page.inner_text('.block h2')
          and not page.is_disabled('#compareBtn'), str(pressed()))
    page.locator('.chip[data-name="氏名"]').first.click()
    check('setup: another column replaces a guessed key', pressed() == ['氏名'] and '推定' not in page.inner_text('.block h2'), str(pressed()))
    page.locator('.chip[data-name="氏名"]').first.click()
    check('setup: key required', page.is_disabled('#compareBtn') and 'キー' in page.inner_text('.actions .hint'))
    page.locator('.chip[data-name="社員番号"]').first.click()
    check('setup: key selected', page.get_attribute('.chip[data-name="社員番号"]', 'aria-pressed') == 'true')
    # Short of memory (simulated): asked before reading; やめる stays on this screen.
    page.route('**/api/memory_check', lambda r: r.fulfill(status=200, content_type='application/json', body=json.dumps(
        {'input_bytes': 2 << 30, 'needed_bytes': 3 << 30, 'available_bytes': 1 << 30, 'enough': False})))
    page.click('#compareBtn')
    page.wait_for_selector('.modal')
    body = page.inner_text('.modal')
    check('setup: a likely memory shortage is told before reading', '3.0\u00a0GB' in body and '1.0\u00a0GB' in body, body[:100])
    page.locator('.modal button', has_text='やめる').click()
    check('setup: declining keeps the files and keys', page.is_visible('main.setup') and page.get_attribute('.chip[data-name="社員番号"]', 'aria-pressed') == 'true')
    page.unroute('**/api/memory_check')
    page.click('#compareBtn')
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    idle(page)


def compare_tests(page: Page, tmp: Path) -> None:
    s = page.evaluate('__cdr.state.summary')
    c = s['counts']
    exp = EXPECTED['counts']
    check('counts match samples/expected.json',
          (c['changed'], c['added'], c['deleted'], c['same'], c['units']) ==
          (exp['changed'], exp['added'], exp['deleted'], exp['same'], exp['all']), str(c))
    check('all rows shown first', total(page) == exp['all'])
    st = sel(page)
    first = api(page, 'view', {'order': 'aligned', 'filter': {'kinds': ['changed', 'added', 'deleted']}})
    first_row = api(page, 'rows', {'view': first['view'], 'start': 0, 'count': 1, 'col_start': 0, 'col_count': 1})['rows'][0]
    check('first difference selected on open', st['unit'] == first_row['unit'], f"{st} vs {first_row['unit']}")
    check('detail shows the selected record', page.locator('#detail .dsum .key').inner_text() == ' / '.join(key_of(page, st['unit'])))

    # Next difference through all differences, then the end.
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('Control+Home')
    wait_sel(page)
    units = []
    diffs = c['changed'] + c['added'] + c['deleted']
    for _ in range(diffs + 1):
        page.keyboard.press('j')
        page.wait_for_timeout(40)
        idle(page)
        units.append(sel(page)['unit'])
    check('J visits every difference once, in order', len(set(units[:diffs])) == diffs and units == sorted(units), str(units))
    check('no wrap after the last difference', units[-1] == units[-2] and 'これより後に差分はありません' in page.inner_text('.toast'))
    page.keyboard.press('k')
    page.wait_for_timeout(40)
    idle(page)
    check('K goes back', sel(page)['unit'] == units[diffs - 2])

    # Scope and kind filters.
    page.click('[data-scope="diff"]')
    idle(page)
    check('差分のみ shows only differences', total(page) == diffs, str(total(page)))
    page.click('.kindbtn[data-kind="changed"]')
    idle(page)
    check('kind filter shows only changes', total(page) == c['changed'], str(total(page)))
    page.click('.kindbtn[data-kind="changed"]')
    idle(page)
    check('kind filter toggles back', total(page) == diffs)
    page.click('[data-scope="all"]')
    idle(page)

    # Search: Enter jumps to the next cell containing the text, Shift+Enter back.
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('Control+Home')
    wait_sel(page)
    page.keyboard.press('Control+f')
    page.keyboard.press('Control+a')
    page.keyboard.type('営業企画課')
    page.keyboard.press('Enter')
    page.wait_for_function("document.querySelector('.searchbox .hits').textContent !== ''")
    idle(page)
    hits_text = page.inner_text('.searchbox .hits')
    st = wait_sel(page)
    first_row = st['rows']['right']
    check('search counts the matching rows', hits_text == '6行', hits_text)
    check('search moves the cursor to the matching cell (right table)',
          page.evaluate('__cdr.state.active') == 'right' and
          '営業企画課' in page.evaluate('''(() => { const p = __cdr.state.panes.right; const r = p.rowsEl.querySelector(`.row[data-row="${p.sel.focus.row}"] .c[data-col="${p.sel.focus.col}"]`); return r ? r.textContent : ''; })()'''))
    check('matching cells are marked', page.locator('#paneR .c.hit').count() >= 1)
    page.keyboard.press('Enter')
    idle(page)
    second = wait_sel(page)['rows']['right']
    page.keyboard.press('Shift+Enter')
    idle(page)
    back = wait_sel(page)['rows']['right']
    check('Enter goes to the next match, Shift+Enter back', second > first_row and back == first_row, f'{first_row} {second} {back}')
    page.locator('#paneR .scroller').focus()
    page.keyboard.press('F3')
    idle(page)
    check('F3 continues the search from the table', wait_sel(page)['rows']['right'] == second)

    # Narrowing to the searched text.
    page.click('.toolbar button:has-text("絞り込む")')
    idle(page)
    n = total(page)
    vis = page.evaluate("__cdr.visible('right')")
    check('絞り込む keeps rows containing the text', n == 6 and all('営業企画課' in ''.join(v['text']) for v in vis if not v['pending']), str(n))
    page.fill('input.search', 'no-such-value-xyz')
    page.press('input.search', 'Enter')
    idle(page)
    check('empty result says what to do', page.locator('#paneL .noresult button', has_text='絞り込みを解除').is_visible())
    page.locator('#paneL .noresult button').click()
    idle(page)
    check('clearing the filter restores all rows', total(page) == exp['all'] and page.get_attribute('.toolbar button:has-text("絞り込む")', 'aria-pressed') == 'false')
    page.fill('input.search', '')
    page.keyboard.press('Escape')

    # Go to key and to a line.
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('Control+g')
    page.wait_for_selector('.popover:not([hidden]) .goform input')
    page.locator('.popover:not([hidden]) .goform input').first.fill('001004')
    page.locator('.popover:not([hidden]) button', has_text='キーで移動').click()
    st = wait_sel(page)
    check('go to key selects the record', key_of(page, st['unit']) == ['001004'])
    page.keyboard.press('Control+g')
    page.wait_for_selector('.popover:not([hidden]) select')
    page.locator('.popover:not([hidden]) select').select_option('right')
    page.locator('.popover:not([hidden]) input[inputmode=numeric]').fill('2')
    page.locator('.popover:not([hidden]) button', has_text='行で移動').click()
    page.wait_for_timeout(100)
    st = wait_sel(page)
    check('go to right line 2 selects that record', key_of(page, st['unit']) == ['001019'], str(key_of(page, st['unit'])))

    # Detail: character-level difference of the selected change.
    page.keyboard.press('Control+g')
    page.locator('.popover:not([hidden]) .goform input').first.fill('001004')
    page.locator('.popover:not([hidden]) button', has_text='キーで移動').click()
    wait_sel(page)
    page.wait_for_selector('#detail .dtable del')
    check('detail marks removed and added characters',
          page.inner_text('#detail .dtable del') == '一' and page.inner_text('#detail .dtable ins') == '企画')
    page.locator('#detail .tabs button', has_text='元の行').click()
    raw = page.inner_text('#detail .rawgrid')
    check('raw tab shows the original text', '"001004"' in raw and '営業企画課' in raw)
    page.locator('#detail .tabs button', has_text='変更した列').click()

    # Original file order keeps the selection and shows each file's lines.
    unit = sel(page)['unit']
    page.click('[data-order="source"]')
    # The selection is located in both files' orders after the views open.
    page.wait_for_function('''() => {
      const s = __cdr.state.sel, r = s && s.rows.right;
      return __cdr.state.panes.right.model.order === 'right_source' && r != null &&
        __cdr.visible('right').some(v => v.row === r && !v.pending);
    }''')
    st = sel(page)
    check('元ファイル順 keeps the selected record', st['unit'] == unit)
    rrow = st['rows'].get('right')
    rvis = [v for v in page.evaluate("__cdr.visible('right')") if v['row'] == rrow]
    check('right pane shows the record at its line in the right file', bool(rvis) and rvis[0]['text'][1] == '53', str(rvis[:1]))
    page.click('[data-order="aligned"]')
    idle(page)

    # A click on a column name selects the column, as in a spreadsheet; with
    # Shift, the columns up to it. Sorting is in the header's menu.
    col = page.evaluate("__cdr.state.panes.left.order.indexOf(__cdr.state.panes.left.model.slots.findIndex(s => s.name_left === '給与'))")
    name_col = page.evaluate("__cdr.state.panes.left.order.indexOf(__cdr.state.panes.left.model.slots.findIndex(s => s.name_left === '氏名'))")
    page.locator('#paneL .hrow .c[data-slot]', has_text='氏名').click()
    r = page.evaluate('__cdr.state.panes.left.sel.range')
    check('a click on a column name selects the column', r == {'r0': 0, 'r1': exp['all'] - 1, 'c0': name_col, 'c1': name_col}
          and page.evaluate('__cdr.state.sort') is None, str(r))
    page.locator('#paneL .hrow .c[data-slot]', has_text='部署').click(modifiers=['Shift'])
    dept_col = page.evaluate("__cdr.state.panes.left.order.indexOf(__cdr.state.panes.left.model.slots.findIndex(s => s.name_left === '部署'))")
    r = page.evaluate('__cdr.state.panes.left.sel.range')
    check('Shift+click on a column name selects the columns up to it', r['c0'] == min(name_col, dept_col)
          and r['c1'] == max(name_col, dept_col) and r['r1'] == exp['all'] - 1, str(r))
    # Sort by a column (header menu) and by the number of changed columns.
    page.locator('#paneL .hrow .c[data-slot]', has_text='給与').click(button='right')
    page.locator('.menu .mi', has_text='昇順で並べ替え').click()
    idle(page)
    hdr = page.locator('#paneL .hrow .c[data-slot]', has_text='給与').inner_text()
    page.evaluate(f"__cdr.jumpTo(0); __cdr.state.panes.left.ensureColVisible({col})")
    idle(page)
    vals = page.evaluate(f'''[...document.querySelectorAll('#paneL .row:not(.pending) .c[data-col="{col}"]')]
        .map(e => e.textContent).filter(Boolean)''')
    check('the header menu sorts ascending with a mark', '▲' in hdr and vals == sorted(vals), f'{hdr} {vals[:5]}')
    page.locator('#paneL .hrow .c[data-slot]', has_text='給与').click(button='right')
    page.locator('.menu .mi', has_text='並べ替えを解除').click()
    idle(page)
    check('the header menu removes the sort', '▲' not in page.locator('#paneL .hrow .c[data-slot]', has_text='給与').inner_text()
          and page.evaluate('__cdr.state.sort') is None)

    # Copy: all rows of the left pane, and a range with column names.
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('Control+a')
    page.keyboard.press('Control+c')
    page.wait_for_function('window.__lastCopy')
    text = page.evaluate('window.__lastCopy')
    lines = text.split('\r\n')
    check('Ctrl+A, Ctrl+C copies every row with the column names', len(lines) == exp['all'] + 2 and lines[0].startswith('社員番号\t') and lines[-1] == '', str(len(lines)))
    check('copied values are the original values', lines[1].split('\t')[0] == '001001')
    page.evaluate('window.__lastCopy = null')
    page.locator('#paneL .row[data-row="0"] .c[data-col="0"]').click()
    page.locator('#paneL .row[data-row="2"] .c[data-col="2"]').click(modifiers=['Shift'])
    page.locator('#paneL .row[data-row="1"] .c[data-col="1"]').click(button='right')
    page.locator('.menu .mi', has_text='列名を付けてコピー').click()
    page.wait_for_function('window.__lastCopy')
    text = page.evaluate('window.__lastCopy')
    check('range copy with column names', text.split('\r\n')[0] == '社員番号\t氏名\t部署' and len(text.split('\r\n')) == 5, repr(text[:80]))
    page.evaluate('window.__lastCopy = null')
    page.locator('#paneL .row[data-row="1"] .c[data-col="1"]').click(button='right')
    page.locator('.menu .mi', has_text='左右を並べてコピー').click()
    page.wait_for_function('window.__lastCopy')
    text = page.evaluate('window.__lastCopy')
    check('side-by-side copy names both sides', text.split('\r\n')[0].startswith('左:社員番号\t左:氏名\t左:部署\t右:社員番号'), repr(text[:80]))

    # The table's top-left corner selects everything; the copy has the column names.
    page.evaluate('window.__lastCopy = null')
    page.locator('#paneR .hrow .c.corner').first.click()
    page.keyboard.press('Control+c')
    page.wait_for_function('window.__lastCopy')
    lines = page.evaluate('window.__lastCopy').split('\r\n')
    check('corner click selects all; copy includes column names',
          lines[0].startswith('社員番号\t') and len(lines) == exp['all'] + 2 and page.evaluate('__cdr.state.active') == 'right',
          f'{lines[0][:40]} {len(lines)}')
    check('the selection size is shown', '選択 ' in page.inner_text('.pwrap[data-side="right"] .csel'))

    # Keyboard cursor and pane switch.
    page.locator('#paneL .row[data-row="3"] .c[data-col="1"]').click()
    page.keyboard.press('ArrowDown')
    page.keyboard.press('ArrowRight')
    f = page.evaluate('__cdr.state.panes.left.sel.focus')
    check('arrow keys move the cell cursor', f == {'row': 4, 'col': 2}, str(f))
    page.keyboard.press('Shift+ArrowDown')
    r = page.evaluate('__cdr.state.panes.left.sel.range')
    check('Shift+arrow extends the selection', r == {'r0': 4, 'r1': 5, 'c0': 2, 'c1': 2}, str(r))
    # Dragging past the bottom edge keeps selecting and scrolls the table.
    page.evaluate('__cdr.jumpTo(0)')
    idle(page)
    box = page.locator('#paneL .scroller').bounding_box()
    cell = page.locator('#paneL .row[data-row="2"] .c[data-col="1"]').bounding_box()
    page.mouse.move(cell['x'] + 5, cell['y'] + 5)
    page.mouse.down()
    page.mouse.move(cell['x'] + 5, box['y'] + box['height'] + 40, steps=5)
    page.wait_for_timeout(600)
    page.mouse.up()
    r = page.evaluate('__cdr.state.panes.left.sel.range')
    check('drag selection scrolls past the edge', r['r0'] == 2 and r['r1'] > page.evaluate('__cdr.state.panes.left.visibleRows') and
          page.evaluate('__cdr.state.panes.left.top') > 0, str(r))
    page.locator('#paneL .row[data-row="3"] .c[data-col="1"]').click() if page.locator('#paneL .row[data-row="3"]').count() else None
    page.keyboard.press('F6')
    check('F6 moves to the other table', page.evaluate("document.activeElement.closest('#paneR') !== null"))
    page.keyboard.press('F6')

    # Column width: drag the header edge; double-click fits the values.
    head = page.locator('#paneL .hrow .c[data-slot]', has_text='氏名')
    w0 = head.bounding_box()['width']
    rz = head.locator('.rz').bounding_box()
    page.mouse.move(rz['x'] + 3, rz['y'] + 5)
    page.mouse.down()
    page.mouse.move(rz['x'] + 83, rz['y'] + 5, steps=4)
    page.mouse.up()
    page.wait_for_timeout(100)
    w1 = head.bounding_box()['width']
    check('dragging a header edge widens the column in both tables',
          abs(w1 - w0 - 80) <= 2 and abs(page.locator('#paneR .hrow .c[data-slot]', has_text='氏名').bounding_box()['width'] - w1) <= 1,
          f'{w0} -> {w1}')
    rz = head.locator('.rz').bounding_box()
    page.mouse.move(rz['x'] + 3, rz['y'] + 5)
    page.mouse.down()
    # Past the minimum width the pointer ends inside the same header cell.
    page.mouse.move(rz['x'] + 3 - (w1 - 32) - 12, rz['y'] + 5, steps=6)
    page.mouse.up()
    page.wait_for_timeout(100)
    check('resizing does not sort', page.evaluate('__cdr.state.sort') is None and abs(head.bounding_box()['width'] - 32) <= 1,
          f"{page.evaluate('__cdr.state.sort')} {head.bounding_box()['width']}")
    head.locator('.rz').dblclick()
    page.wait_for_timeout(100)
    check('double-click fits the column to its values', head.bounding_box()['width'] < w1, str(head.bounding_box()['width']))

    # Hide a column (display only) and show it again.
    page.locator('#paneL .hrow .c[data-slot]', has_text='役職').click(button='right')
    page.locator('.menu .mi', has_text='この列を隠す').click()
    page.wait_for_timeout(100)
    check('hidden column leaves the table', page.locator('#paneL .hrow .c[data-slot]', has_text='役職').count() == 0)
    check('hidden columns are indicated', page.locator('button.note', has_text='隠した列 1').is_visible())
    check('hiding does not change the counts', page.evaluate('__cdr.state.summary.counts.changed') == c['changed'])
    page.locator('button.note').click()
    page.wait_for_timeout(100)
    check('hidden columns come back', page.locator('#paneL .hrow .c[data-slot]', has_text='役職').count() == 1)

    # A record whose only change is in a hidden column still shows as changed (AC-08).
    full = api(page, 'groups', {'rule': 'full'})
    only = next(g['first_unit'] for g in full['items'] if [c['name'] for c in g['changes']] == ['給与'])
    page.locator('#paneL .hrow .c[data-slot]', has_text='給与').click(button='right')
    page.locator('.menu .mi', has_text='この列を隠す').click()
    page.wait_for_timeout(100)
    row = page.locator(f'#paneL .row[data-unit="{only}"]')
    if row.count() == 0:
        page.evaluate('(u) => __cdr.state.locateAndSelect(u, __cdr.state.panes.left)', only)
        idle(page)
    row.locator('.c.num').click()
    wait_sel(page)
    page.wait_for_function("document.querySelector('#detail .dtable') && document.querySelector('#detail .dtable').textContent.includes('給与')")
    check('a change only in a hidden column keeps the row marked changed', row.locator('.c.mk').inner_text().strip() == '~',
          row.locator('.c.mk').inner_text())
    check('the detail shows the change in the hidden column', '給与' in page.inner_text('#detail .dtable'))
    page.locator('button.note').click()
    page.wait_for_timeout(100)

    # Changed columns only.
    page.locator('.toolbar button', has_text='設定').click()
    page.locator('.popover:not([hidden]) label', has_text='変更のある列だけ表示').click()
    page.wait_for_timeout(100)
    names = page.evaluate("[...document.querySelectorAll('#paneL .hrow .c[data-slot]')].map(e => e.textContent)")
    changed_names = {sl['name_left'] for sl in page.evaluate('__cdr.state.panes.left.model.slots')
                     if sl['key'] or (sl['pair'] is not None and s['pair_changed'][sl['pair']] > 0)}
    check('changed-columns-only shows keys and changed columns', set(names) <= changed_names, f'{names}')
    page.locator('.popover:not([hidden]) label', has_text='変更のある列だけ表示').click()
    page.keyboard.press('Escape')

    # Export the differences to CSV.
    page.click('[data-scope="diff"]')
    idle(page)
    out = tmp / 'export.csv'
    answers.append(str(out))
    page.locator('.toolbar button', has_text='書き出す').click()
    page.locator('.popover:not([hidden]) input[name="xformat"][value="csv"]').check()
    page.locator('.popover:not([hidden]) select[aria-label="文字コード"]').select_option('utf8_bom')
    page.locator('.popover:not([hidden]) button', has_text='保存先を選んで書き出す').click()
    page.wait_for_function("document.querySelector('.toast') && document.querySelector('.toast').textContent.includes('書き出しました')")
    data = out.read_bytes()
    lines = data[3:].decode('utf-8').split('\r\n')
    check('export writes the shown rows as CSV with BOM',
          data.startswith(b'\xef\xbb\xbf') and lines[0].startswith('種類,左の行,右の行,変更した列,左:社員番号') and len(lines) == diffs + 2,
          f'{lines[0][:60]} {len(lines)}')
    page.evaluate("document.querySelector('.toast').hidden = true")
    out_sj = tmp / 'export_sjis.csv'
    answers.append(str(out_sj))
    page.locator('.toolbar button', has_text='書き出す').click()
    page.locator('.popover:not([hidden]) select[aria-label="文字コード"]').select_option('windows31j')
    page.locator('.popover:not([hidden]) button', has_text='保存先を選んで書き出す').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('書き出しました')")
    check('export in Shift_JIS', out_sj.read_bytes().decode('cp932').startswith('種類,左の行,右の行,変更した列,左:社員番号'))
    page.evaluate("document.querySelector('.toast').hidden = true")
    out_x = tmp / 'export.xlsx'
    answers.append(str(out_x))
    page.locator('.toolbar button', has_text='書き出す').click()
    check('the dialog remembers the last choices', page.locator('.popover:not([hidden]) input[name="xformat"][value="csv"]').is_checked())
    guard = page.locator('.popover:not([hidden]) label', has_text='式として読まれる値')
    check('the formula guard is offered for CSV, on by default', guard.is_visible() and guard.locator('input').is_checked())
    page.locator('.popover:not([hidden]) input[name="xformat"][value="xlsx"]').check()
    check('the encoding and formula guard are hidden for Excel',
          not page.locator('.popover:not([hidden]) select[aria-label="文字コード"]').is_visible() and not guard.is_visible())
    page.locator('.popover:not([hidden]) button', has_text='保存先を選んで書き出す').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('書き出しました')")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(out_x)
        ws = wb['比較結果']
        fills = {ws.cell(r, c).fill.fgColor.rgb for r in range(2, ws.max_row + 1) for c in range(1, ws.max_column + 1)
                 if ws.cell(r, c).fill.fill_type}
        check('Excel export marks changes, additions and deletions',
              ws.max_row == diffs + 1 and {'FFFDEEB0', 'FFD9F2E1', 'FFFBE0DE'} <= fills and '比較条件' in wb.sheetnames and
              ws.cell(2, 5).value == ws.cell(2, 5).value.strip() and ws.cell(2, 5).data_type == 's', f'{ws.max_row} {fills}')
    except ImportError:
        check('Excel export written', out_x.read_bytes().startswith(b'PK'))
    page.evaluate("document.querySelector('.toast').hidden = true")
    answers.append(str(SAMPLES / 'before.csv'))
    page.locator('.toolbar button', has_text='書き出す').click()
    page.locator('.popover:not([hidden]) button', has_text='保存先を選んで書き出す').click()
    page.wait_for_function("document.querySelector('.toast.error')")
    check('export refuses to overwrite a compared file', '書き出せません' in page.inner_text('.toast'))
    page.keyboard.press('Escape')
    page.click('[data-scope="all"]')
    idle(page)

    # Layouts.
    page.click('[data-layout="stack"]')
    idle(page)
    l = page.locator('#paneL').bounding_box()
    r = page.locator('#paneR').bounding_box()
    check('上下 stacks the tables', r['y'] >= l['y'] + l['height'] - 2 and abs(l['x'] - r['x']) < 2)
    page.click('[data-layout="side"]')
    idle(page)
    l = page.locator('#paneL').bounding_box()
    r = page.locator('#paneR').bounding_box()
    check('左右 puts them side by side', r['x'] >= l['x'] + l['width'] - 2 and abs(l['y'] - r['y']) < 2)
    rows_open = page.evaluate("__cdr.state.panes.left.visibleRows")
    page.click('#detailToggle')
    page.wait_for_timeout(100)
    rows_closed = page.evaluate("__cdr.state.panes.left.visibleRows")
    check('hiding the detail gives the tables more rows', rows_closed > rows_open, f'{rows_open} -> {rows_closed}')
    page.click('#detailToggle')


def failure_tests(page: Page, url: str, tmp: Path) -> None:
    # Keys that cannot tell records apart, seen in the files before comparing.
    dup = tmp / 'dup.csv'
    dup.write_text('id,name\n1,a\n2,b\n2,c\n4,d\n', encoding='utf-8')
    answers[:] = [str(dup), str(dup)]
    page.goto(url)
    page.wait_for_selector('main.setup')
    page.locator('.card').nth(0).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .pv')
    page.locator('.card').nth(1).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    check('a column with repeated values is not guessed as the key',
          page.get_attribute('.chip[data-name="name"]', 'aria-pressed') == 'true')
    choose_key(page, 'id')
    problem = page.inner_text('.block .error')
    check('a key with repeated values is told before comparing, with its lines',
          '「id」の値「2」が、左のファイルの3行目と4行目' in problem and page.is_disabled('#compareBtn'), problem)
    # Repeated only in combination: found by the comparison, which lists them.
    combo = tmp / 'combo.csv'
    combo.write_text('id,name\n1,a\n2,b\n2,b\n3,\n', encoding='utf-8')
    answers[:] = [str(combo), str(combo)]
    page.locator('.card').nth(0).locator('button', has_text='別のファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .pv')
    page.locator('.card').nth(1).locator('button', has_text='別のファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    page.wait_for_function("!!document.querySelector('.chip[data-name=\"name\"]')")
    choose_key(page, 'id')
    page.locator('.chip[data-name="name"]').first.click()
    problem = page.inner_text('.block .error')
    check('an empty value in a key column is told before comparing', '「name」の値が、左のファイルの5行目で空です' in problem, problem)
    combo.write_text('id,name\n1,a\n2,b\n2,b\n', encoding='utf-8')
    answers[:] = [str(combo), str(combo)]
    page.locator('.card').nth(0).locator('button', has_text='別のファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .pv')
    page.locator('.card').nth(1).locator('button', has_text='別のファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    page.wait_for_function("!document.getElementById('compareBtn').disabled")
    page.click('#compareBtn')
    page.wait_for_selector('main.failed')
    t = page.inner_text('main.failed')
    check('duplicate keys stop the comparison with their lines', '同じキーのレコード 1組（2件）' in t and '3、4' in t, t[:200])
    check('the reason comes first and how to put it right last',
          page.inner_text('main.failed .reason').startswith('キーの値が同じレコード') and '直し方' in t
          and 'キーの列を増やして' in page.inner_text('main.failed .fix'))
    page.locator('button', has_text='条件を変更する').click()
    page.wait_for_selector('main.setup .pv')
    check('back to the setup with the same files and key',
          page.get_attribute('.chip[data-name="id"]', 'aria-pressed') == 'true')
    # A fresh start lists the comparison; one click loads its files and key.
    api(page, 'reset')
    page.reload()
    page.wait_for_selector('main.setup')
    page.wait_for_selector('main.setup .recent li')
    check('recent comparisons are listed', 'combo.csv と combo.csv' in page.inner_text('main.setup .recent'))
    page.locator('.recentbtn').first.click()
    page.wait_for_selector('.card[data-side=right] .pv')
    page.wait_for_selector('.chip[data-name="id"][aria-pressed="true"]')
    check('a recent comparison loads its files and key', page.locator('main.setup .recent').is_hidden())
    page.locator('main.setup .actions button', has_text='設定').click()
    page.locator('.popover:not([hidden]) label', has_text='最近の比較を記録する').click()
    page.wait_for_function("JSON.parse(localStorage.getItem('cdr-settings') || '{}').keepHistory === false")
    saved = page.evaluate("JSON.parse(localStorage.getItem('cdr-settings'))")
    check('turning history off forgets the recent comparisons', saved['recent'] == [] and 'lastKeys' not in saved)

    # The licenses of the software built into the app.
    page.locator('.popover:not([hidden]) button', has_text='ライセンス情報').click()
    page.wait_for_selector('.modal .mtext')
    text = page.inner_text('.modal .mtext')
    check('the license texts can be shown', 'tauri' in text and 'Permission is hereby granted' in text and len(text) > 100000, str(len(text)))
    page.keyboard.press('Escape')
    check('the license dialog closes with Escape', page.locator('.modal').count() == 0)


def review_summary(page: Page) -> dict:
    return api(page, 'review_summary')


def review_tests(page: Page, url: str, tmp: Path) -> None:
    page.goto(url)
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    idle(page)
    check('review states are off by default', page.locator('#paneL .hrow .c.rv').count() == 0)
    page.locator('.toolbar button', has_text='設定').click()
    page.locator('.popover:not([hidden]) label', has_text='確認状況とメモを使う').click()
    page.keyboard.press('Escape')
    page.wait_for_selector('#paneL .hrow .c.rv')
    check('turning reviews on adds the review column', page.locator('#paneL .hrow .c.rv').count() == 1)
    page.wait_for_function("document.querySelector('.sidebar .strip').textContent.includes('確認 0')")
    page.locator('.sidebar .strip').click()
    page.wait_for_selector('.sidebar .rcount')
    check('the side panel shows the progress', '確認済み 0／18件' in page.inner_text('.sidebar .rcount'))

    # F1 lists the keyboard operations, including the review keys while reviews are on.
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('F1')
    page.wait_for_selector('.modal .keytable')
    keys = page.inner_text('.modal')
    check('F1 lists the keyboard operations', 'J ／ K' in keys and 'Space' in keys and 'Ctrl+Shift+C' in keys, keys[:80])
    page.keyboard.press('Escape')
    check('the key list closes and the table keeps the focus', page.locator('.modal').count() == 0 and page.evaluate("document.activeElement.classList.contains('scroller')"))

    # Space marks the selected record and moves to the next unchecked one.
    page.locator('#paneL .scroller').focus()
    first = wait_sel(page)['unit']
    page.keyboard.press('Space')
    page.wait_for_function('(u) => __cdr.state.sel && __cdr.state.sel.unit !== u && __cdr.state.sel.unit >= 0', arg=first)
    idle(page)
    s = review_summary(page)
    check('Space marks the record checked and moves on', s['checked'] == 1 and wait_sel(page)['unit'] != first)
    page.wait_for_function("document.querySelector('.sidebar .rcount').textContent.includes('1／18')")
    rv = page.locator(f'#paneL .row[data-unit="{first}"] .c.rv')
    check('the review column shows the state', rv.inner_text() == '✓')

    # 3 marks for investigation; M writes a note; Ctrl+Z takes the note back.
    second = wait_sel(page)['unit']
    page.keyboard.press('3')
    page.wait_for_function('(u) => __cdr.state.sel.unit !== u', arg=second)
    page.keyboard.press('Shift+n')
    page.keyboard.press('Shift+n')
    idle(page)
    page.locator(f'#paneL .row[data-unit="{second}"] .c[data-col="1"]').click()
    wait_sel(page)
    page.keyboard.press('m')
    page.keyboard.type('仕入先に確認する')
    page.keyboard.press('Escape')
    page.wait_for_function("document.querySelector('.sidebar').textContent.includes('メモあり 1件')")
    d = api(page, 'detail', {'unit': second})
    check('3 and M set 要調査 and a note', d['review']['status'] == 'investigate' and d['review']['memo'] == '仕入先に確認する')
    page.locator('#paneL .scroller').focus()
    page.keyboard.press('Control+z')
    page.wait_for_function("!document.querySelector('.sidebar').textContent.includes('メモあり')")
    check('Ctrl+Z undoes the last change', api(page, 'detail', {'unit': second})['review']['memo'] == '')

    # Several rows: the change is shown with its counts first, then applied and undoable.
    page.click('[data-scope="diff"]')
    idle(page)
    page.locator('#paneL .row[data-row="0"] .c[data-col="1"]').click()
    page.locator('#paneL .row[data-row="5"] .c[data-col="1"]').click(modifiers=['Shift'])
    page.keyboard.press('Space')
    page.wait_for_selector('.modal')
    body = page.inner_text('.modal')
    check('a bulk change shows what it applies to', '6件' in body and '内訳' in body and '変わった列' in body, body[:120])
    page.locator('.modal button.primary').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('確認済みにしました')")
    after = review_summary(page)
    check('the bulk change is applied to those records', after['checked'] >= 6, str(after))
    page.locator('.toast .tact').click()
    page.wait_for_function("document.querySelector('.toast').textContent.includes('元に戻しました')")
    check('元に戻す in the notice undoes it', review_summary(page)['checked'] == 1)
    page.locator('#paneL .row[data-row="1"] .c[data-col="1"]').click()
    page.keyboard.press('Shift+Space')
    idle(page)
    check('Shift+Space selects the row and marks nothing', page.evaluate('__cdr.state.panes.left.sel.whole') == 'rows'
          and page.locator('.modal').count() == 0 and review_summary(page)['checked'] == 1)

    # Change patterns: the rule is shown; a group filters the tables; full patterns can be marked at once.
    page.locator('.sidebar select[aria-label="まとめ方"]').select_option('full')
    page.wait_for_selector('.sidebar .gitem')
    check('the active rule is explained', 'すべての列と、その前後の値が同じ' in page.inner_text('.sidebar .gexplain'))
    page.locator('.sidebar .gitem').first.click()
    idle(page)
    check('a pattern shows only its records', total(page) == 5 and page.locator('.toolbar .chips button', has_text='パターンで絞り込み中').is_visible())
    page.locator('.sidebar .gmark').click()
    page.wait_for_selector('.modal')
    check('the pattern bulk change names the change', '部署（5件）' in page.inner_text('.modal'))
    page.locator('.modal button.primary').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('5件を確認済みにしました')")
    page.locator('.toolbar .chips button', has_text='パターンで絞り込み中').click()
    idle(page)
    # A column left out of the patterns: shown as the applied condition, and counted apart.
    page.locator('.sidebar .gskip summary').click()
    page.locator('.sidebar .gskip label', has_text='部署').click()
    page.wait_for_function("document.querySelector('.sidebar .ginfo').textContent.includes('まとめに使わない列だけが変わったもの 5件')")
    check('a skipped column is shown as the applied condition', page.inner_text('.sidebar .gskip summary') == 'まとめに使わない列：部署')
    glist = page.inner_text('.sidebar .glist')
    check('patterns are formed without the skipped column', '部署' not in glist and '役職' in glist, glist[:120])
    page.locator('.sidebar .gskip label', has_text='部署').click()
    page.wait_for_function("!document.querySelector('.sidebar .ginfo').textContent.includes('まとめに使わない列だけ')")
    check('clearing the skipped column restores the patterns', page.inner_text('.sidebar .gskip summary') == 'まとめに使わない列：なし')
    page.locator('.sidebar select[aria-label="まとめ方"]').select_option('per_column')
    page.wait_for_selector('.sidebar .gnote')
    check('per-column groups say when records changed elsewhere too', 'ほかの列も変わっています' in page.inner_text('.sidebar .glist'))

    # Review filter.
    page.locator('.sidebar .rfilters button', has_text='未確認').click()
    idle(page)
    s = review_summary(page)
    check('the unchecked filter shows the unchecked differences', total(page) == s['unchecked'], f"{total(page)} {s['unchecked']}")
    page.locator('.sidebar .rfilters button', has_text='未確認').click()
    idle(page)

    # Export with the review columns.
    out = tmp / 'review.csv'
    answers.append(str(out))
    page.locator('.toolbar button', has_text='書き出す').click()
    page.locator('.popover:not([hidden]) input[name="xformat"][value="csv"]').check()
    page.locator('.popover:not([hidden]) button', has_text='保存先を選んで書き出す').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('書き出しました')")
    head = out.read_bytes().decode('utf-8-sig').split('\r\n')[0]
    check('exports include the review state and memo', '確認状況,メモ' in head, head[:80])

    # Clearing every state of the comparison is confirmed first and can be undone.
    before = review_summary(page)
    page.locator('.sidebar .rclear').click()
    page.wait_for_selector('.modal')
    check('clearing all states says how many', f"{before['checked']}件" in page.inner_text('.modal'), page.inner_text('.modal')[:120])
    page.locator('.modal button.primary').click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('未確認に戻しました')")
    check('clearing all states leaves every difference unchecked', review_summary(page)['unchecked'] == before['total'])
    page.locator('.toast .tact').click()
    page.wait_for_function("document.querySelector('.toast').textContent.includes('元に戻しました')")
    after = review_summary(page)
    check('clearing all states can be undone', (after['checked'], after['investigate']) == (before['checked'], before['investigate']))
    # The harness is stopped without a last write (the app writes on exit): wait for the background write.
    for _ in range(50):
        if not review_summary(page)['unsaved']:
            break
        page.wait_for_timeout(200)
    saved_summary.update(review_summary(page))


saved_summary: dict = {}
errors_expected: list[str] = []


def review_resume_test(page: Page, url: str) -> None:
    page.goto(url)
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    page.wait_for_function("document.querySelector('.toast') && document.querySelector('.toast').textContent.includes('前回の確認状況を読み込みました')", timeout=30000)
    s = review_summary(page)
    want = saved_summary
    check('the same files and conditions bring the review states back',
          s['restored'] == want['checked'] + want['investigate'] and (s['checked'], s['investigate']) == (want['checked'], want['investigate']),
          f'{s} vs {want}')


def save_error_test(page: Page, url: str, blocked: Path) -> None:
    """The review folder cannot be written: the failure is shown, and saving again works once it can."""
    page.goto(url)
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    idle(page)
    page.wait_for_function("document.querySelector('.sidebar .strip').textContent.includes('確認 0')")
    page.locator('#paneL .scroller').focus()
    wait_sel(page)
    page.keyboard.press('Space')
    chip = page.locator('.toolbar .chips button', has_text='確認状況を保存できません')
    chip.wait_for(timeout=10000)
    check('a failed save is shown in the toolbar', chip.is_visible() and '保存先のフォルダ' in (chip.get_attribute('title') or ''))
    check('the change stays in the app', review_summary(page)['checked'] == 1)
    # Going back to the setup screen now would lose the change: asked first.
    page.locator('.toolbar button', has_text='条件を変更').click()
    page.wait_for_selector('.modal')
    check('leaving with unsaved review states asks first', '保存していない確認状況は失われます' in page.inner_text('.modal'))
    page.locator('.modal button', has_text='やめる').click()
    check('declining keeps the comparison open', page.locator('main.compare').count() == 1)
    blocked.unlink()
    blocked.mkdir()
    chip.click()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('確認状況を保存しました')")
    check('saving again works once the folder can be written', chip.count() == 0 and any(blocked.glob('*.json')))


def low_memory_test(page: Page, url: str, tmp: Path) -> None:
    """Memory running low during a comparison: it stops with a reason, and the files can be compared again."""
    page.goto(url)
    page.wait_for_selector('main.setup')
    for i, f in enumerate(('before.csv', 'after.csv')):
        answers.append(str(SAMPLES / f))
        page.locator('.card').nth(i).locator('button', has_text='ファイルを選ぶ').click()
        page.wait_for_selector(f'.card[data-side={("left", "right")[i]}] .pv')
    choose_key(page, '社員番号')
    page.click('#compareBtn')
    page.wait_for_selector('main.failed', timeout=30000)
    t = page.inner_text('main.failed')
    check('low memory stops the comparison with a reason', 'メモリが足りなくなったため' in t and 'ほかのアプリを閉じて' in t, t[:120])
    page.locator('button', has_text='条件を変更する').click()
    page.wait_for_selector('main.setup .pv')
    check('after a low-memory stop, the same files are kept for another try', page.locator('main.setup .pv').count() == 2)


def formats_test(page: Page, url: str, tmp: Path) -> None:
    """Other layouts: Excel's "Unicode text" (UTF-16, tabs) without quotes
    on the left, fixed-width Shift_JIS on the right."""
    left = tmp / 'unicode.txt'
    rows = [['番号', '氏名', '部署'], ['001', '山田 太郎', '営業部'], ['002', '佐藤 花子', '人事部'],
            ['003', '鈴木 一郎', '開発部'], ['004', '5" 画面', '総務部']]
    left.write_bytes(b'\xff\xfe' + '\r\n'.join('\t'.join(r) for r in rows).encode('utf-16-le') + '\r\n'.encode('utf-16-le'))
    right = tmp / 'fixed.dat'
    fixed = [['番号', '氏名', '部署'], ['001', '山田 太郎', '営業部'], ['002', '佐藤 花子', '経理部'],
             ['003', '鈴木 一郎', '開発部'], ['005', '田中 次郎', '総務部']]
    pad = lambda v, w: v.encode('cp932') + b' ' * (w - len(v.encode('cp932')))
    right.write_bytes(b''.join(pad(r[0], 4) + pad(r[1], 10) + pad(r[2], 6) + b'\r\n' for r in fixed))

    page.goto(url)
    page.wait_for_selector('main.setup')
    # A fixed-width file whose first try as CSV fails on a quote: the card says
    # the delimiter is missing, not the quote error, and shows the lines.
    quoted = tmp / 'quote.dat'
    quoted.write_bytes('番号  品名\r\n001   5" 画面, 黒\r\n'.encode('cp932'))
    answers.append(str(quoted))
    page.locator('.card').nth(0).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .rawlines')
    check('formats: a file without the delimiter is told so, with its lines',
          'カンマがありません' in page.inner_text('.card[data-side=left] .error') and page.locator('.card[data-side=left] .rawline').count() == 2)
    answers.append(str(left))
    page.locator('.card').nth(0).locator('button', has_text='別のファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=left] .pv')
    card = page.locator('.card[data-side=left]')
    enc = card.locator('select[aria-label="文字コード"]').input_value()
    delim = card.locator('select[aria-label="区切り"]').input_value()
    quoted = card.locator('label.check', has_text='引用符').locator('input').is_checked()
    check('formats: UTF-16 with tabs is recognised', (enc, delim) == ('utf16le', '9'), f'{enc} {delim}')
    check('formats: a file that does not use quotes is read without them', not quoted and '5" 画面' in card.locator('.pv').inner_text())

    answers.append(str(right))
    page.locator('.card').nth(1).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    rc = page.locator('.card[data-side=right]')
    check('formats: Shift_JIS is suggested for the fixed-width file', rc.locator('select[aria-label="文字コード"]').input_value() == 'windows31j')
    rc.locator('select[aria-label="区切り"]').select_option('fixed')
    rc.locator('input[aria-label="列の幅"]').wait_for()
    check('formats: fixed width waits for the widths', page.is_disabled('#compareBtn') and '列の幅' in page.inner_text('.actions .hint')
          and rc.locator('.rawline').count() == 5)
    check('formats: bytes are the unit for Shift_JIS', rc.locator('select[aria-label="幅の単位"]').input_value() == 'bytes')
    # Changing the encoding while the widths are still to be entered keeps fixed width.
    rc.locator('select[aria-label="文字コード"]').select_option('euc_jp')
    page.wait_for_function("document.querySelector('.card[data-side=right] select[aria-label=\"文字コード\"]').value === 'euc_jp'")
    check('formats: another encoding keeps the layout being set', rc.locator('input[aria-label="列の幅"]').count() == 1)
    rc.locator('select[aria-label="文字コード"]').select_option('windows31j')
    page.wait_for_function("document.querySelector('.card[data-side=right] select[aria-label=\"文字コード\"]').value === 'windows31j'")
    # A width too large to be one: told, and the settings stay.
    rc.locator('input[aria-label="列の幅"]').fill('4, 99999999999')
    rc.locator('input[aria-label="列の幅"]').press('Enter')
    page.wait_for_function("document.querySelector('.card[data-side=right] .error') && document.querySelector('.card[data-side=right] .error').textContent.includes('列の幅を読み取れません')")
    check('formats: an impossible width is told and the settings stay',
          rc.locator('input[aria-label="列の幅"]').input_value() == '4, 99999999999' and page.is_disabled('#compareBtn')
          and '列の幅を直して' in page.inner_text('.actions .hint'))
    rc.locator('input[aria-label="列の幅"]').fill('４，10, 6')
    rc.locator('input[aria-label="列の幅"]').press('Enter')
    page.wait_for_function("document.querySelectorAll('.card[data-side=right] .pv th').length === 3")
    heads = rc.locator('.pv th').all_inner_texts()
    check('formats: the widths cut the columns (full-width digits accepted)', heads == ['番号', '氏名', '部署'], str(heads))
    check('formats: the cuts are shown on the lines', rc.locator('.rawline').first.locator('.piece').count() == 3)
    check('formats: the focus stays in the widths after Enter', page.evaluate("document.activeElement.getAttribute('aria-label')") == '列の幅')
    choose_key(page, '番号')
    # Widths typed but not applied yet: Ctrl+Enter applies them first and does not compare with the old ones.
    rc.locator('input[aria-label="列の幅"]').fill('4, 10, 4, 2')
    rc.locator('input[aria-label="列の幅"]').press('Control+Enter')
    page.wait_for_selector('.card[data-side=right] .rawline .piece:nth-child(4)')
    check('formats: Ctrl+Enter applies typed widths instead of comparing with the old ones',
          page.is_visible('main.setup') and api(page, 'status')['phase']['state'] == 'idle')
    # The fourth column cut there has no name: told before comparing.
    check('formats: a column without a name is told before comparing',
          '右のファイルの4列目は、列名が空です' in page.inner_text('.block .error') and page.is_disabled('#compareBtn'))
    check('formats: the focus stays in the widths after Ctrl+Enter', page.evaluate("document.activeElement.getAttribute('aria-label')") == '列の幅')
    rc.locator('input[aria-label="列の幅"]').fill('4, 10, 6')
    rc.locator('input[aria-label="列の幅"]').press('Enter')
    page.wait_for_function("document.querySelector('.card[data-side=right] .rawline .piece:nth-child(3)') && document.querySelector('.card[data-side=right] .rawline .piece:nth-child(3)').textContent.startsWith('部署')")
    # A delimiter typed with the IME on, as a full-width character, is used as half-width.
    card.locator('select[aria-label="区切り"]').select_option('other')
    card.locator('input[aria-label="区切り文字"]').fill('\uff1a')
    page.wait_for_function("document.querySelectorAll('.card[data-side=left] .pv th').length === 1")
    check('formats: a full-width delimiter is taken as half-width',
          card.locator('input[aria-label="区切り文字"]').input_value() == ':')
    card.locator('select[aria-label="区切り"]').select_option('9')
    page.wait_for_function("document.querySelectorAll('.card[data-side=left] .pv th').length === 3")
    check('formats: a key whose column went away for a moment comes back with it',
          page.get_attribute('.chip[data-name="番号"]', 'aria-pressed') == 'true')
    page.click('#compareBtn')
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    s = page.evaluate('__cdr.state.summary')
    c = s['counts']
    check('formats: the two layouts compare by value', (c['same'], c['changed'], c['added'], c['deleted']) == (2, 1, 1, 1), str(c))
    check('formats: how each file was read is recorded',
          s['left']['layout'] == '区切り タブ、引用符なし' and s['right']['layout'].startswith('固定長（バイト単位の幅 4, 10, 6'),
          f"{s['left']['layout']} / {s['right']['layout']}")


def headerless_test(page: Page, url: str, tmp: Path) -> None:
    """Files without a line of column names, compared by column position; a
    file chosen can be taken off again."""
    old, new = tmp / 'nohead_old.csv', tmp / 'nohead_new.csv'
    old.write_text('001,山田,2026-04-01,300000\n002,佐藤,2026-04-01,310000\n', encoding='utf-8')
    new.write_text('001,山田,2026-04-01,300000\n002,佐藤,2026-05-01,320000\n003,鈴木,2026-05-01,280000\n', encoding='utf-8')
    page.goto(url)
    page.wait_for_selector('main.setup')
    for i, f in enumerate((old, new)):
        answers.append(str(f))
        page.locator('.card').nth(i).locator('button', has_text='ファイルを選ぶ').click()
        page.wait_for_selector(f'.card[data-side={("left", "right")[i]}] .pv')
    card = page.locator('.card[data-side=left]')
    check('headerless: a first line of data is not taken for column names',
          not card.locator('label.check', has_text='1行目は列名').locator('input').is_checked()
          and card.locator('.pv th').all_inner_texts() == ['列1', '列2', '列3', '列4'])
    check('headerless: the key is guessed among the numbered columns',
          page.get_attribute('.chip[data-name="列1"]', 'aria-pressed') == 'true' and not page.is_disabled('#compareBtn'))
    # The × takes a file off; the side is as before choosing.
    page.locator('.card[data-side=right] button[aria-label="右のファイルの選択を解除"]').click()
    check('a file chosen can be taken off', page.locator('.card[data-side=right] .dropzone').count() == 1
          and page.is_disabled('#compareBtn') and page.evaluate("document.activeElement.textContent") == 'ファイルを選ぶ')
    answers.append(str(new))
    page.locator('.card').nth(1).locator('button', has_text='ファイルを選ぶ').click()
    page.wait_for_selector('.card[data-side=right] .pv')
    page.wait_for_function("!document.getElementById('compareBtn').disabled")
    page.click('#compareBtn')
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    c = page.evaluate('__cdr.state.summary.counts')
    check('headerless: compared by column position', (c['same'], c['changed'], c['added'], c['deleted']) == (1, 1, 1, 0), str(c))
    kinds = [b.get_attribute('aria-label') for b in page.locator('.kindbtn').all()]
    check('the kinds read 左のみ and 右のみ', kinds == ['変更 1件', '左のみ 0件', '右のみ 1件', '一致 1件'], str(kinds))


def spreadsheet_keys_test(page: Page, url: str, tmp: Path) -> None:
    """Keys as in a spreadsheet: Ctrl+arrows stop where values start and
    end, Shift+Space and Ctrl+Space select rows and columns."""
    rows = 'id,a,b\n1,x,\n2,y,\n3,,p\n4,,\n5,z,q\n6,w,r\n'
    old, new = tmp / 'keys_old.csv', tmp / 'keys_new.csv'
    old.write_text(rows, encoding='utf-8')
    new.write_text(rows.replace('6,w,r', '6,w,s'), encoding='utf-8')
    page.goto(url)
    page.wait_for_selector('main.setup')
    for i, f in enumerate((old, new)):
        answers.append(str(f))
        page.locator('.card').nth(i).locator('button', has_text='ファイルを選ぶ').click()
        page.wait_for_selector(f'.card[data-side={("left", "right")[i]}] .pv')
    page.wait_for_function("!document.getElementById('compareBtn').disabled")
    page.click('#compareBtn')
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    idle(page)

    def at(row: int, col: int, side: str = 'left') -> bool:
        try:
            page.wait_for_function(f"(f => f.row === {row} && f.col === {col})(__cdr.state.panes.{side}.sel.focus)", timeout=3000)
            return True
        except Exception:
            return False

    def rng() -> dict:
        return page.evaluate('__cdr.state.panes.left.sel.range')

    page.locator('#paneL .row[data-row="0"] .c[data-col="1"]').click()
    page.keyboard.press('Control+ArrowDown')
    check('keys: Ctrl+Down stops at the last value before an empty cell', at(1, 1))
    page.keyboard.press('Control+ArrowDown')
    check('keys: Ctrl+Down goes over empty cells to the next value', at(4, 1))
    page.keyboard.press('Control+ArrowDown')
    page.keyboard.press('Control+ArrowUp')
    check('keys: Ctrl+Up comes back along the values', at(4, 1))
    page.keyboard.press('Control+Shift+ArrowUp')
    check('keys: Ctrl+Shift+Up selects up to the edge', at(1, 1) and rng() == {'r0': 1, 'r1': 4, 'c0': 1, 'c1': 1}, str(rng()))
    page.locator('#paneL .row[data-row="0"] .c[data-col="0"]').click()
    page.keyboard.press('Control+ArrowRight')
    check('keys: Ctrl+Right stops before an empty cell', at(0, 1))
    page.keyboard.press('Control+ArrowRight')
    check('keys: Ctrl+Right with nothing after goes to the last column', at(0, 2))
    page.locator('#paneL .row[data-row="0"] .c[data-col="1"]').click()
    page.keyboard.press('Control+ArrowDown')
    page.keyboard.press('Control+ArrowDown')
    check('keys: Ctrl+Down pressed quickly goes on from where the last one ends', at(4, 1))

    page.keyboard.press('Shift+ArrowDown')
    page.keyboard.press('Shift+Space')
    check('keys: Shift+Space selects the rows, the cursor stays',
          rng() == {'r0': 4, 'r1': 5, 'c0': 0, 'c1': 2} and at(5, 1), str(rng()))
    page.locator('#paneL .row[data-row="2"] .c[data-col="2"]').click()
    page.keyboard.press('Control+Space')
    check('keys: Ctrl+Space selects the column, the cursor stays', rng() == {'r0': 0, 'r1': 5, 'c0': 2, 'c1': 2} and at(2, 2), str(rng()))
    page.evaluate('window.__lastCopy = null')
    page.keyboard.press('Control+c')
    page.wait_for_function('window.__lastCopy')
    check('keys: a column selected is copied with its name', page.evaluate('window.__lastCopy').split('\r\n')[:4] == ['b', '', '', 'p'],
          repr(page.evaluate('window.__lastCopy')[:40]))
    page.keyboard.press('Shift+Space')
    check('keys: rows and columns together are the whole table', rng() == {'r0': 0, 'r1': 5, 'c0': 0, 'c1': 2}, str(rng()))
    page.keyboard.press('ArrowDown')
    check('keys: an arrow after a whole selection moves from the cursor', at(3, 2) and rng() == {'r0': 3, 'r1': 3, 'c0': 2, 'c1': 2})
    page.keyboard.press('Control+Shift+Space')
    check('keys: Ctrl+Shift+Space selects everything', rng() == {'r0': 0, 'r1': 5, 'c0': 0, 'c1': 2} and at(3, 2))

    page.keyboard.press('Control+End')
    check('keys: Ctrl+End goes to the last cell', at(5, 2))
    page.keyboard.press('Control+Home')
    check('keys: Ctrl+Home goes to the first cell', at(0, 0))
    page.keyboard.press('Enter')
    page.keyboard.press('Enter')
    page.keyboard.press('Shift+Enter')
    check('keys: Enter and Shift+Enter move down and up', at(1, 0))
    page.keyboard.press('Control+PageDown')
    check('keys: Ctrl+PageDown goes to the right table', page.evaluate("document.activeElement.closest('#paneR') !== null"))
    page.keyboard.press('Control+PageUp')
    check('keys: Ctrl+PageUp goes back to the left table', page.evaluate("document.activeElement.closest('#paneL') !== null"))
    # The line number selects the row; with Shift, the rows up to it.
    page.locator('#paneL .row[data-row="1"] .c.num').click()
    page.locator('#paneL .row[data-row="3"] .c.num').click(modifiers=['Shift'])
    check('keys: line numbers select whole rows', rng() == {'r0': 1, 'r1': 3, 'c0': 0, 'c1': 2}, str(rng()))
    page.keyboard.press('F1')
    page.wait_for_selector('.keytable')
    check('keys: the key list names the new keys', all(k in page.inner_text('.keytable') for k in ('Ctrl+矢印キー', 'Shift+Space', 'Ctrl+Space')))
    page.keyboard.press('Escape')

    # Enter works a focused button, not the table.
    before = page.evaluate('__cdr.state.panes.left.sel.focus')
    page.locator('.toolbar button', has_text='設定').focus()
    page.keyboard.press('Enter')
    check('keys: Enter on a focused button presses it', page.locator('.popover:not([hidden])').count() == 1
          and page.evaluate('__cdr.state.panes.left.sel.focus') == before)
    page.keyboard.press('Escape')

    # Keys pressed while a Ctrl+arrow is worked out (slowed down here) come after it, in turn.
    page.evaluate('''() => { const f = window.fetch; window.fetch = async (u, o) => {
        if (String(u).includes('api/edge')) await new Promise((r) => setTimeout(r, 300)); return f(u, o); }; }''')
    page.locator('#paneL .row[data-row="0"] .c[data-col="1"]').click()
    page.keyboard.press('Control+ArrowDown')
    page.keyboard.press('ArrowDown')
    page.keyboard.press('Control+ArrowDown')
    page.keyboard.press('ArrowRight')
    check('keys: keys pressed during a slow Ctrl+arrow follow it in turn', at(4, 2), str(page.evaluate('__cdr.state.panes.left.sel.focus')))
    page.locator('#paneL .row[data-row="0"] .c[data-col="1"]').click()
    page.keyboard.press('Control+ArrowDown')
    page.locator('#paneL .row[data-row="3"] .c[data-col="2"]').click()
    page.wait_for_timeout(500)
    check('keys: a click while a Ctrl+arrow is worked out wins over it', at(3, 2))
    # A table without a cursor starts from the record selected.
    page.locator('#paneL .row[data-row="4"] .c[data-col="1"]').click()
    page.keyboard.press('Control+PageDown')
    page.keyboard.press('Shift+Space')
    r = page.evaluate('__cdr.state.panes.right.sel.range')
    check('keys: in the other table, Shift+Space selects the row of the record selected', r['r0'] == r['r1'] == 4
          and page.evaluate('__cdr.state.sel.rows.right') == 4, str(r))


def key_guess_and_vanished_file_test(page: Page, url: str, tmp: Path) -> None:
    """An id named like 社員ID is guessed before a name; a file gone between
    choosing and comparing is told in one piece, with how to put it right."""
    old, new = tmp / 'staff_old.csv', tmp / 'staff_new.csv'
    old.write_text('氏名,社員ID,部署\n山田,A1,営業\n佐藤,A2,人事\n', encoding='utf-8')
    new.write_text('氏名,社員ID,部署\n山田,A1,営業\n佐藤,A2,総務\n', encoding='utf-8')
    page.goto(url)
    page.wait_for_selector('main.setup')
    for i, f in enumerate((old, new)):
        answers.append(str(f))
        page.locator('.card').nth(i).locator('button', has_text='ファイルを選ぶ').click()
        page.wait_for_selector(f'.card[data-side={("left", "right")[i]}] .pv')
    check('a column named like 社員ID is guessed as the key before 氏名',
          page.get_attribute('.chip[data-name="社員ID"]', 'aria-pressed') == 'true'
          and page.get_attribute('.chip[data-name="氏名"]', 'aria-pressed') == 'false')
    new.unlink()
    page.click('#compareBtn')
    page.wait_for_selector('main.failed', timeout=30000)
    reason, fix = page.inner_text('main.failed .reason'), page.inner_text('main.failed .fix')
    check('a file gone before comparing: the reason is not cut inside its brackets',
          reason.endswith('）。') and 'ファイルを読めません（' in reason and fix == 'パスとアクセス権を確認してください。', f'{reason!r} / {fix!r}')


def startup_test(page: Page, url: str, tmp: Path) -> None:
    """Files and keys given at start (dropped on the exe, 送る, command line) open the comparison directly."""
    from urllib.parse import quote
    files = '&'.join(f"open={quote(str(SAMPLES / f))}" for f in ('before.csv', 'after.csv'))
    page.goto(f"{url}&{files}&key={quote('社員番号')}")
    page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=30000)
    check('files and a key given at start open the comparison', page.evaluate('__cdr.state.summary.counts.changed') == 12)
    # The version is shown in the settings; a bug in the screen's code is told, not silent.
    page.locator('.toolbar button', has_text='設定').click()
    page.wait_for_function("[...document.querySelectorAll('.popover:not([hidden]) .muted')].some(e => e.textContent.startsWith('Recolign '))")
    check('the settings show the version', True)
    page.keyboard.press('Escape')
    page.evaluate("setTimeout(() => { throw new Error('a test of the screen error net') })")
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('画面の内部エラー')")
    check('an unexpected error in the screen is told once', True)
    errors_expected.append('a test of the screen error net')
    api(page, 'reset')  # a new start of the app has no comparison yet
    page.goto(f"{url}&{files}&key={quote('ない列')}")
    page.reload()
    page.wait_for_selector('main.setup .pv')
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('比較を始めませんでした')")
    check('a key missing from the files stops on the setup screen with the reason', page.locator('main.setup .card .pv').count() == 2)
    api(page, 'reset')
    files = '&'.join(f"open={quote(str(SAMPLES / f))}" for f in ('ない.csv', 'after.csv'))
    page.goto(f"{url}&{files}&key={quote('社員番号')}")
    page.reload()
    page.wait_for_function("document.querySelector('.toast:not([hidden])') && document.querySelector('.toast').textContent.includes('左のファイルを読み込めないため')")
    check('a file that cannot be read stops the start with the reason',
          page.locator('main.setup .card[data-side=left] .error').count() == 1 and page.is_disabled('#compareBtn'))


def screenshots(browser, url: str, out: Path) -> list[dict]:
    out.mkdir(parents=True, exist_ok=True)
    res = []
    for w, h in ((1366, 768), (1920, 1080)):
        for scale in (1.0, 1.25, 1.5):
            ctx = browser.new_context(viewport={'width': round(w / scale), 'height': round(h / scale)},
                                      device_scale_factor=scale, bypass_csp=True)
            page = ctx.new_page()
            page.goto(url)
            page.wait_for_function('window.__cdr && window.__cdr.ready', timeout=60000)
            idle(page)
            name = f'{w}x{h}@{int(scale * 100)}.png'
            page.screenshot(path=str(out / name))
            res.append({'screen': f'{w}x{h}', 'scale': scale, 'file': name,
                        'rows_visible': page.evaluate("__cdr.state.panes.left.visibleRows"),
                        'toolbar_height': page.evaluate("document.querySelector('.toolbar').getBoundingClientRect().height"),
                        'no_page_horizontal_scroll': page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')})
            ctx.close()
    return res


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--harness', type=Path, required=True)
    ap.add_argument('--screens', type=Path)
    a = ap.parse_args()
    errors: list[str] = []
    requests: list[str] = []
    with sync_playwright() as pw, tempfile.TemporaryDirectory() as tmp:
        browser = pw.chromium.launch(executable_path=CHROMIUM if Path(CHROMIUM).exists() else None, args=['--no-sandbox'])
        sample_args = ['--left', str(SAMPLES / 'before.csv'), '--right', str(SAMPLES / 'after.csv'), '--key', '社員番号',
                       '--data-dir', str(Path(tmp) / 'reviews')]
        blocked = Path(tmp) / 'blocked'
        blocked.write_text('not a folder')
        blocked_args = ['--left', str(SAMPLES / 'before.csv'), '--right', str(SAMPLES / 'after.csv'), '--key', '社員番号',
                        '--data-dir', str(blocked)]
        for name, fn, args in (
            ('compare', None, []),
            ('failure', failure_tests, []),
            ('review', review_tests, sample_args),
            ('resume', review_resume_test, sample_args),
            ('save error', save_error_test, blocked_args),
            ('low memory', low_memory_test, ['--memory-floor', str(2 ** 64 - 1)]),
            ('startup', startup_test, []),
            ('formats', formats_test, []),
            ('headerless', headerless_test, []),
            ('keys', spreadsheet_keys_test, []),
            ('key guess', key_guess_and_vanished_file_test, []),
        ):
            h = Harness(a.harness, args)
            ctx = browser.new_context(viewport={'width': 1366, 'height': 768}, bypass_csp=True)
            page = ctx.new_page()
            page.on('pageerror', lambda e: errors.append(str(e)))
            page.on('request', lambda r: requests.append(r.url))
            try:
                page.on('dialog', on_dialog)
                if name == 'compare':
                    setup_flow(page, h.url)
                    compare_tests(page, Path(tmp))
                elif name == 'resume':
                    # Settings live in the browser profile of each context: turn reviews on again.
                    page.add_init_script("localStorage.setItem('cdr-settings', JSON.stringify({reviewEnabled: true}))")
                    fn(page, h.url)
                elif name == 'save error':
                    page.add_init_script("localStorage.setItem('cdr-settings', JSON.stringify({reviewEnabled: true}))")
                    fn(page, h.url, blocked)
                else:
                    fn(page, h.url, Path(tmp))
            except Exception:
                traceback.print_exc()
                check(f'{name}: ran to the end', False)
            finally:
                ctx.close()
                h.close()
        if a.screens:
            h = Harness(a.harness, ['--left', str(SAMPLES / 'before.csv'), '--right', str(SAMPLES / 'after.csv'), '--key', '社員番号'])
            try:
                shots = screenshots(browser, h.url, a.screens)
                (a.screens / 'screens.json').write_text(json.dumps(shots, ensure_ascii=False, indent=1), encoding='utf-8')
                for s in shots:
                    print(s)
            finally:
                h.close()
        browser.close()
    errors = [e for e in errors if not any(x in e for x in errors_expected)]
    check('no script errors on the page', not errors, '; '.join(errors[:3]))
    outside = [u for u in requests if not u.startswith(('http://127.0.0.1:', 'data:', 'blob:'))]
    check('the screens request nothing outside this machine (AC-16)', not outside and len(requests) > 100, '; '.join(outside[:3]))
    failed = [r for r in results if not r[1]]
    print(f'\n{len(results) - len(failed)}/{len(results)} passed')
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
